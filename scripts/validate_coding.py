#!/usr/bin/env python3
"""
validate_coding.py — checks that a coded-data table is internally consistent
and traceable to its sources. Standard library only; openpyxl is used if
present so .xlsx workbooks can be read directly, otherwise export the sheets
to CSV first. Paths below are relative to the plugin root.

    python3 scripts/validate_coding.py \
        --coded    coded_data.csv   (or the coded-data .xlsx workbook) \
        --codebook codebook.csv     (or the codebook .xlsx) \
        --manifest manifest.csv     (optional; defaults to the "manifest" sheet of an .xlsx --coded) \
        --sources  ./sources        (directory holding the files named in the manifest) \
        --summary  frequencies.csv  (optional; regenerated and compared) \
        [--codebook-version v2]     (optional; every row must carry it)

Checks
  1. every Text excerpt occurs verbatim in its source file. "Verbatim" means:
     runs of whitespace (including non-breaking spaces and line breaks) are
     collapsed to one space; everything else — case, punctuation, quotation
     marks — must match exactly. Rows with Language = translated cannot be
     checked and make the result INCOMPLETE
  2. every document (or document part) the manifest marks as coded has rows,
     and vice versa; manifest Units matches the number of distinct Unit IDs
  3. every Code ID exists in the codebook; inactive codes are flagged
  4. Assignment IDs are unique; no (Document, Unit, Code) is assigned twice
  5. Prompting values are in the allowed set (blank is a warning)
  6. frequencies regenerated from the rows match the summary table, if given;
     a summary with a "Source type" column is compared per stratum
Document parts collapse to their parent for document counts: the manifest's
"Parent document" column when present, else the DOC.§n / DOC§n convention.
Sheet names in .xlsx are matched case-insensitively; a single-sheet workbook
is accepted for any table.

Result and exit status
  PASS        (0) every check ran and none failed
  FAIL        (1) at least one check failed
  INCOMPLETE  (2) no check failed, but a required check could not run — no
                  manifest, no --sources, no --summary, a document without a
                  Source file, or translated rows. Lines marked SKIP say which.
                  INCOMPLETE is not a pass: do not report it as one.
Warnings (WARN) do not change the result. The script never modifies any file.
"""
import argparse, csv, os, re, sys
from collections import defaultdict

PROMPTING_VALUES = {"prompted", "volunteered", "mixed", "unclear", "n/a"}
UNCODED = "UNCODED"
REQ_CODED = ["Document ID", "Unit ID", "Assignment ID", "Text excerpt", "Code ID", "Prompting", "Codebook version", "Language"]
REQ_CODEBOOK = ["Code ID"]
REQ_MANIFEST = ["Document ID", "Source file", "Status"]
REQ_SUMMARY = ["Code ID"]

# ---------------------------------------------------------------- helpers
def key(s):
    return re.sub(r"[\s_/]+", "", s or "").lower()

def read_table(path, sheet=None):
    """Rows as dicts with normalized keys. For .xlsx, `sheet` is matched
    case-insensitively; a single-sheet workbook is accepted; otherwise a missing
    sheet returns None (never a silent fallback to a different sheet)."""
    if path.lower().endswith(".xlsx"):
        try:
            import openpyxl  # optional
        except ImportError:
            sys.exit(f"{path}: reading .xlsx needs openpyxl (pip install openpyxl) — or export the sheet to CSV")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet:
            match = [n for n in wb.sheetnames if n.lower() == sheet.lower()]
            if match:
                ws = wb[match[0]]
            elif len(wb.sheetnames) == 1:
                ws = wb[wb.sheetnames[0]]
            else:
                return None
        else:
            ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        head = [key(str(h)) if h is not None else "" for h in rows[0]]
        return [dict(zip(head, ["" if v is None else str(v).strip() for v in r]))
                for r in rows[1:] if any(v not in (None, "") for v in r)]
    with open(path, newline="", encoding="utf-8-sig") as f:
        rd = csv.DictReader(f)
        return [{key(k): (v or "").strip() for k, v in row.items() if k is not None} for row in rd]

def col(row, *names):
    for n in names:
        v = row.get(key(n))
        if v is not None:
            return v.strip()
    return ""

def missing_columns(rows, required):
    if not rows:
        return []
    have = set(rows[0].keys())
    return [c for c in required if key(c) not in have]

def squash(s):
    """Whitespace-normalize only: runs of whitespace (incl. NBSP, line breaks)
    become one space. Case and punctuation are left as they are."""
    return re.sub(r"\s+", " ", s.replace("\u00a0", " ")).strip()

def parent_doc(did, parents=None):
    if parents and parents.get(did):
        return parents[did]
    return re.split(r"\.?§", did, maxsplit=1)[0]

class Report:
    def __init__(self): self.fail, self.warn, self.ok, self.skip = [], [], [], []
    def f(self, m): self.fail.append(m)
    def w(self, m): self.warn.append(m)
    def o(self, m): self.ok.append(m)
    def s(self, m): self.skip.append(m)   # a required check that could not run

# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--coded", required=True)
    ap.add_argument("--codebook", required=True)
    ap.add_argument("--manifest")
    ap.add_argument("--sources")
    ap.add_argument("--summary")
    ap.add_argument("--codebook-version")
    a = ap.parse_args()
    R = Report()

    coded = read_table(a.coded, "coded_data" if a.coded.lower().endswith(".xlsx") else None)
    if coded is None:
        sys.exit(f"{a.coded}: no sheet named 'coded_data'")
    if not coded:
        sys.exit("coded table is empty")
    codebook = read_table(a.codebook, "codebook" if a.codebook.lower().endswith(".xlsx") else None)
    if codebook is None:
        sys.exit(f"{a.codebook}: no sheet named 'codebook'")

    if a.manifest:
        manifest = read_table(a.manifest, "manifest" if a.manifest.lower().endswith(".xlsx") else None)
        if manifest is None:
            R.f(f"{a.manifest}: no sheet named 'manifest'"); manifest = []
    elif a.coded.lower().endswith(".xlsx"):
        manifest = read_table(a.coded, "manifest")
        if manifest is None:
            R.s("coded workbook has no 'manifest' sheet — coverage and excerpts not checked"); manifest = []
    else:
        manifest = []
    summary = []
    if a.summary:
        summary = read_table(a.summary, "summary" if a.summary.lower().endswith(".xlsx") else None)
        if summary is None:
            R.f(f"{a.summary}: no sheet named 'summary'"); summary = []

    # ---- required columns (a missing column is a FAIL, never a silent pass)
    for name, rows, req in (("coded table", coded, REQ_CODED), ("codebook", codebook, REQ_CODEBOOK),
                            ("manifest", manifest, REQ_MANIFEST), ("summary", summary, REQ_SUMMARY)):
        if rows:
            miss = missing_columns(rows, req)
            if miss:
                R.f(f"{name} is missing required column(s): {', '.join(miss)}")
    if any("missing required column" in m and m.startswith("coded table") for m in R.fail):
        finish(R, coded, {}, {}); return

    # ---- 3. codes exist in codebook
    cb = {}
    for r in codebook:
        cid = col(r, "Code ID")
        if cid:
            cb[cid] = (col(r, "Status") or "active").lower()
    used = defaultdict(int)
    for r in coded:
        used[col(r, "Code ID")] += 1
    if used.get("", 0):
        R.f(f"{used['']} rows have an empty Code ID (use {UNCODED} for units that fit no code)")
    code_fail = False
    for cid, n in sorted(used.items()):
        if cid in ("", UNCODED):
            continue
        if cid not in cb:
            R.f(f"code {cid!r} used {n}× but not in codebook"); code_fail = True
        elif cb[cid] != "active":
            R.w(f"code {cid!r} used {n}× but marked {cb[cid]!r} in codebook")
    real_codes = [c for c in used if c not in ("", UNCODED)]
    if not code_fail and not used.get("", 0):
        R.o(f"codes: all {len(real_codes)} code IDs present in codebook")

    # ---- codebook version stamp
    if a.codebook_version:
        bad = [r for r in coded if col(r, "Codebook version") != a.codebook_version]
        if bad:
            R.f(f"{len(bad)} rows not coded under {a.codebook_version} (found: {sorted({col(r, 'Codebook version') for r in bad})})")
        else:
            R.o(f"version: all rows coded under {a.codebook_version}")

    # ---- 4. uniqueness
    seen_a, seen_t, dup = set(), set(), False
    for r in coded:
        aid = col(r, "Assignment ID")
        if not aid:
            R.f(f"row for {col(r,'Document ID')}/{col(r,'Unit ID')} has an empty Assignment ID"); dup = True
        elif aid in seen_a:
            R.f(f"duplicate Assignment ID {aid!r}"); dup = True
        seen_a.add(aid)
        t = (col(r, "Document ID"), col(r, "Unit ID"), col(r, "Code ID"))
        if t in seen_t:
            R.f(f"{aid or '(no id)'}: code {t[2]!r} assigned twice to unit {t[0]}/{t[1]}"); dup = True
        seen_t.add(t)
    if not dup:
        R.o("uniqueness: assignment IDs and (document, unit, code) triples are unique")

    # ---- 5. prompting values
    pfail, pblank = 0, 0
    for r in coded:
        v = col(r, "Prompting")
        if v == "":
            pblank += 1
        elif v not in PROMPTING_VALUES:
            pfail += 1
            R.f(f"{col(r,'Assignment ID')}: Prompting {v!r} not in {sorted(PROMPTING_VALUES)}")
    if pblank:
        R.w(f"{pblank} rows have a blank Prompting value (use n/a for non-elicited text)")
    if not pfail:
        R.o("prompting: all values allowed")

    # ---- 2. coverage against manifest
    docs_in_coded = defaultdict(set)
    for r in coded:
        docs_in_coded[col(r, "Document ID")].add(col(r, "Unit ID"))
    srcfile, parents = {}, {}
    if manifest:
        cov_fail = False
        for m in manifest:
            did = col(m, "Document ID"); status = col(m, "Status").lower()
            srcfile[did] = col(m, "Source file"); parents[did] = col(m, "Parent document")
            units = col(m, "Units")
            if status in ("coded", "done", "complete", "completed"):
                if did not in docs_in_coded:
                    R.f(f"manifest says {did} is coded but it has no rows"); cov_fail = True
                elif units == "":
                    R.w(f"{did}: manifest Units is blank — unit count not checked")
                elif not units.isdigit():
                    R.w(f"{did}: manifest Units {units!r} is not an integer — unit count not checked")
                elif int(units) != len(docs_in_coded[did]):
                    R.f(f"{did}: manifest says {units} units, coded table has {len(docs_in_coded[did])} distinct Unit IDs"); cov_fail = True
            elif did in docs_in_coded:
                R.w(f"{did} has rows but manifest status is {status or 'blank'!r}")
        for did in docs_in_coded:
            if did not in srcfile:
                R.f(f"{did} has rows but is not in the manifest"); cov_fail = True
        if not cov_fail:
            R.o(f"coverage: {len(docs_in_coded)} documents/parts, manifest and coded table agree")
    elif a.manifest:
        R.f("manifest is empty (header only, or no rows)")
    else:
        R.s("no manifest — coverage not checked")

    # ---- 1. excerpts in sources
    if a.sources and not manifest:
        R.s("--sources given but there is no manifest to map documents to files — excerpts not checked")
    elif a.sources:
        cache, checked, missing, skipped, filefail, nofile = {}, 0, 0, 0, 0, {}
        for r in coded:
            did = col(r, "Document ID"); ex = col(r, "Text excerpt"); aid = col(r, "Assignment ID")
            if not ex:
                R.f(f"{aid}: empty Text excerpt"); missing += 1; continue
            if col(r, "Language").lower() == "translated":
                skipped += 1; continue
            fn = srcfile.get(did)
            if not fn:
                nofile[did] = nofile.get(did, 0) + 1; continue
            p = os.path.join(a.sources, fn)
            if p not in cache:
                try:
                    with open(p, encoding="utf-8", errors="replace") as f:
                        cache[p] = squash(f.read())
                except OSError:
                    cache[p] = None; R.f(f"source file not found: {p}"); filefail += 1
            if cache[p] is None:
                continue
            checked += 1
            if squash(ex) not in cache[p]:
                missing += 1
                R.f(f"{aid}: excerpt not found verbatim (case- and punctuation-exact, whitespace collapsed) in {fn}: {ex[:60]!r}…")
        for did, n in sorted(nofile.items()):
            R.s(f"{did}: {n} rows not checked against a source — manifest has no Source file for it")
        if skipped:
            R.s(f"{skipped} rows with Language = translated not checked — translated excerpts cannot be matched against the source")
        if checked and not missing and not filefail and not nofile:
            R.o(f"excerpts: all {checked} excerpts found verbatim in their source files")
    else:
        R.s("no --sources given — excerpts not checked against source files")

    # ---- 6. frequencies (pooled, or per source type if the summary has that column)
    stratified = bool(summary) and key("Source type") in summary[0]
    per = defaultdict(lambda: {"assignments": 0, "units": set(), "documents": set()})
    for r in coded:
        cid = col(r, "Code ID")
        if cid in ("", UNCODED):
            continue
        k = (cid, col(r, "Source type")) if stratified else (cid, "")
        d = per[k]; d["assignments"] += 1
        d["units"].add((col(r, "Document ID"), col(r, "Unit ID")))
        d["documents"].add(parent_doc(col(r, "Document ID"), parents))
    if summary:
        mism = 0; listed = set(); uncompared = 0
        for s in summary:
            k = (col(s, "Code ID"), col(s, "Source type") if stratified else "")
            listed.add(k)
            label = f"{k[0]}" + (f" [{k[1]}]" if stratified else "")
            if k not in per:
                R.f(f"summary lists {label} which has no assignments"); mism += 1; continue
            got = per[k]
            for name, alts, val in (("Assignments", ("Assignments", "Count"), got["assignments"]),
                                    ("Units", ("Units",), len(got["units"])),
                                    ("Documents", ("Documents",), len(got["documents"]))):
                want = col(s, *alts)
                if want == "":
                    R.s(f"{label}: summary has no {name} value — not compared"); uncompared += 1
                elif not want.isdigit():
                    R.f(f"{label}: summary {name}={want!r} is not an integer"); mism += 1
                elif int(want) != val:
                    R.f(f"{label}: summary {name}={want}, regenerated {val}"); mism += 1
        for k in per:
            if k not in listed:
                R.f(f"{k[0]}{' ['+k[1]+']' if stratified else ''} has assignments but is missing from summary"); mism += 1
        if not mism and uncompared:
            R.s(f"frequencies: no mismatch found, but {uncompared} count(s) were absent from the summary and not compared")
        elif not mism:
            R.o(f"frequencies: summary matches regenerated counts for {len(per)} code{' × source-type' if stratified else ''} rows")
    else:
        R.s("no --summary given — frequencies regenerated but not compared:")
        for k in sorted(per):
            d = per[k]; R.s(f"    {k[0]}: {d['assignments']} assignments, {len(d['units'])} units, {len(d['documents'])} documents")

    finish(R, coded, docs_in_coded, per)

def finish(R, coded, docs, per):
    n_unc = sum(1 for r in coded if col(r, "Code ID") == UNCODED)
    print(f"rows: {len(coded)}   documents/parts: {len(docs)}   code rows: {len(per)}   UNCODED rows: {n_unc}")
    for m in R.ok:   print("PASS  " + m)
    for m in R.warn: print("WARN  " + m)
    for m in R.skip: print("SKIP  " + m)
    for m in R.fail: print("FAIL  " + m)
    if R.fail:
        print(f"RESULT: FAIL ({len(R.fail)} problem{'s' if len(R.fail) != 1 else ''})"); sys.exit(1)
    if R.skip:
        n = sum(1 for m in R.skip if not m.startswith("    "))
        print(f"RESULT: INCOMPLETE ({n} required check{'s' if n != 1 else ''} did not run — see SKIP lines; not a pass)"); sys.exit(2)
    print("RESULT: PASS"); sys.exit(0)

if __name__ == "__main__":
    main()
